#!/usr/bin/env python3
"""
Auditoría de conciliación para PlantillaFinanzas2026.xlsx (no versionado).

Verifica el invariante del libro: para cada cuenta listada en la hoja `Saldos`,
    C (calculado) = D (saldo inicial) + Ingresos - Egresos + E (rendimientos)
                    + transferencias_entrantes - transferencias_salientes
    G = C - F  debe ser 0,  donde F es el saldo real tecleado desde la app.

Un G distinto de cero es siempre un error de registro:
    G > 0  -> el libro cree que hay más dinero del que existe
              (falta un egreso, o hay un ingreso/transferencia duplicado)
    G < 0  -> el libro cree que hay menos dinero del que existe
              (falta un ingreso o rendimiento, o hay un egreso duplicado)

Solo lectura: nunca escribe en el libro.
Uso: python3 audit_finanzas.py
"""

from openpyxl import load_workbook
from pathlib import Path

EXCEL_PATH = Path.home() / "Codes/mi-criterio/proyectos/P4/PlantillaFinanzas2026.xlsx"

HOJA_SALDOS = "Saldos"
HOJA_TRANS = "Transacciones"
FILA_ENCABEZADO = 29          # los encabezados C/D/E/F/G viven aquí
TOLERANCIA = 0.01             # por debajo de esto es residuo de punto flotante

# Columnas de Transacciones
COL_MONTO, COL_TIPO, COL_CAT, COL_CUENTA, COL_FECHA, COL_DESC = "A", "B", "C", "D", "E", "F"


def cargar():
    """Devuelve (cuentas, movimientos). Los nombres de cuenta salen del libro,
    nunca están escritos aquí — así el script no expone dónde se guarda el dinero."""
    wb = load_workbook(EXCEL_PATH, data_only=True)
    saldos, trans = wb[HOJA_SALDOS], wb[HOJA_TRANS]

    cuentas = []
    for fila in range(FILA_ENCABEZADO + 1, saldos.max_row + 1):
        nombre = saldos[f"B{fila}"].value
        real = saldos[f"F{fila}"].value
        if not isinstance(nombre, str) or not nombre.strip():
            continue
        if not isinstance(real, (int, float)):
            continue          # filas de notas o leyenda, no cuentas
        cuentas.append({
            "fila": fila,
            "nombre": nombre.strip(),
            "inicial": saldos[f"D{fila}"].value or 0,
            "rendim": saldos[f"E{fila}"].value or 0,
            "real": real,
        })

    movs = []
    for fila in range(2, trans.max_row + 1):
        monto = trans[f"{COL_MONTO}{fila}"].value
        cuenta = trans[f"{COL_CUENTA}{fila}"].value
        if monto is None or cuenta is None:
            continue
        movs.append({
            "fila": fila,
            "monto": monto,
            "tipo": trans[f"{COL_TIPO}{fila}"].value,
            "categoria": trans[f"{COL_CAT}{fila}"].value,
            "cuenta": str(cuenta).strip(),
            "fecha": trans[f"{COL_FECHA}{fila}"].value,
            "desc": trans[f"{COL_DESC}{fila}"].value,
        })

    wb.close()
    return cuentas, movs


def conciliar(cuenta, movs):
    n = cuenta["nombre"]

    # Ingreso/Egreso: coincidencia exacta, para que una cuenta no absorba
    # los movimientos de sus sub-cuentas (p. ej. "X" vs "X Plus")
    propios = [m for m in movs if m["cuenta"] == n]
    ing = sum(m["monto"] for m in propios if m["tipo"] == "Ingreso")
    egr = sum(m["monto"] for m in propios if m["tipo"] == "Egreso")

    # Transferencias: la cuenta se codifica como "Origen->Destino"
    trans = [m for m in movs if m["tipo"] == "Transferencia" and "->" in m["cuenta"]]
    entra = sum(m["monto"] for m in trans if m["cuenta"].split("->")[-1].strip() == n)
    sale = sum(m["monto"] for m in trans if m["cuenta"].split("->")[0].strip() == n)

    calc = cuenta["inicial"] + ing - egr + cuenta["rendim"] + entra - sale
    return {**cuenta, "ing": ing, "egr": egr, "entra": entra, "sale": sale,
            "calc": calc, "dif": calc - cuenta["real"],
            "n_movs": len(propios) + sum(1 for m in trans
                                         if n in (m["cuenta"].split("->")[0].strip(),
                                                  m["cuenta"].split("->")[-1].strip()))}


def malformados(movs):
    """Transferencias sin '->': los SUMIFS con comodín las ignoran en silencio."""
    return [m for m in movs if m["tipo"] == "Transferencia" and "->" not in m["cuenta"]]


def huerfanas(movs, cuentas):
    """Transferencias que apuntan a una cuenta que no existe en `Saldos`."""
    conocidas = {c["nombre"] for c in cuentas}
    fuera = []
    for m in movs:
        if m["tipo"] != "Transferencia" or "->" not in m["cuenta"]:
            continue
        origen, destino = [x.strip() for x in m["cuenta"].split("->", 1)]
        desconocidas = {x for x in (origen, destino) if x not in conocidas}
        if desconocidas:
            fuera.append((m, sorted(desconocidas)))
    return fuera


def main():
    cuentas, movs = cargar()
    print(f"{len(movs)} movimientos · {len(cuentas)} cuentas\n")

    ancho = max(len(c["nombre"]) for c in cuentas)
    print(f"{'Cuenta':<{ancho}} {'Calculado':>12} {'Real':>12} {'Dif':>10}  Estado")
    print("-" * (ancho + 46))

    descuadres, total = [], 0.0
    for c in cuentas:
        r = conciliar(c, movs)
        total += r["real"]
        if abs(r["dif"]) <= TOLERANCIA:
            estado = "cuadra"
        elif r["dif"] > 0:
            estado = "SOBRA en el libro (falta egreso / ingreso duplicado)"
            descuadres.append(r)
        else:
            estado = "FALTA en el libro (falta ingreso / egreso duplicado)"
            descuadres.append(r)
        print(f"{r['nombre']:<{ancho}} {r['calc']:>12.2f} {r['real']:>12.2f} "
              f"{r['dif']:>10.2f}  {estado}")

    print("-" * (ancho + 46))
    print(f"{'DINERO TOTAL':<{ancho}} {'':>12} {total:>12.2f}\n")

    mal = malformados(movs)
    if mal:
        print(f"[!] {len(mal)} transferencia(s) sin '->' — los SUMIFS las ignoran:")
        for m in mal:
            print(f"    fila {m['fila']}: {m['monto']} · '{m['cuenta']}'")
        print()

    fuera = huerfanas(movs, cuentas)
    if fuera:
        print(f"[!] {len(fuera)} transferencia(s) hacia cuentas no dadas de alta:")
        for m, desc in fuera:
            print(f"    fila {m['fila']}: {m['monto']} · {m['cuenta']} → {', '.join(desc)}")
        print()

    if not descuadres and not mal and not fuera:
        print("Todo cuadra.")


if __name__ == "__main__":
    main()
